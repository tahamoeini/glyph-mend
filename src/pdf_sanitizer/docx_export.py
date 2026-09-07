from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


_PAGE_MARKER_RE = re.compile(r"^\s*<!--\s*page:\s*(\d+)\s*-->\s*$", re.IGNORECASE)
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.+)$")
_ORDERED_RE = re.compile(r"^(\s*)\d+[.)]\s+(.+)$")
_BLOCKQUOTE_RE = re.compile(r"^\s*>\s?(.*)$")
_FENCE_RE = re.compile(r"^\s*(```+|~~~+)\s*([^\s`]*)\s*$")
_TABLE_SEPARATOR_CELL_RE = re.compile(r"^:?-{3,}:?$")
_INLINE_RE = re.compile(
    r"(\[[^\]]+\]\([^\s)]+\)|\*\*[^*]+\*\*|__[^_]+__|`[^`]+`|(?<!\*)\*[^*\n]+\*(?!\*)|(?<!_)_[^_\n]+_(?!_))"
)


def _style_exists(document: Document, name: str) -> bool:
    try:
        document.styles[name]
        return True
    except KeyError:
        return False


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_inline(paragraph, text: str) -> None:
    position = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("["):
            label_end = token.find("](")
            label = token[1:label_end]
            url = token[label_end + 2 : -1]
            _add_hyperlink(paragraph, label, url)
        elif token.startswith(("**", "__")):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        elif token.startswith(("*", "_")):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _split_table_row(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|"):
        value = value[:-1]

    cells: list[str] = []
    current: list[str] = []
    escaped = False
    in_code = False
    for char in value:
        if escaped:
            current.append(char)
            escaped = False
            continue
        if char == "\\":
            escaped = True
            current.append(char)
            continue
        if char == "`":
            in_code = not in_code
            current.append(char)
            continue
        if char == "|" and not in_code:
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    cells.append("".join(current).strip())
    return cells


def _is_table_separator(line: str) -> bool:
    cells = _split_table_row(line)
    return bool(cells) and all(_TABLE_SEPARATOR_CELL_RE.fullmatch(cell.replace(" ", "")) for cell in cells)


def _looks_like_table(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    if "|" not in lines[index] or "|" not in lines[index + 1]:
        return False
    return _is_table_separator(lines[index + 1])


def _consume_paragraph(lines: list[str], start: int) -> tuple[str, int]:
    values: list[str] = []
    index = start
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            break
        if index != start and (
            _PAGE_MARKER_RE.match(stripped)
            or _HEADING_RE.match(line)
            or _BULLET_RE.match(line)
            or _ORDERED_RE.match(line)
            or _BLOCKQUOTE_RE.match(line)
            or _FENCE_RE.match(line)
            or stripped == "$$"
            or _looks_like_table(lines, index)
        ):
            break
        values.append(stripped)
        index += 1
    return " ".join(values).strip(), index


def _add_code_block(document: Document, lines: Iterable[str], language: str = "") -> None:
    paragraph = document.add_paragraph()
    if _style_exists(document, "No Spacing"):
        paragraph.style = "No Spacing"
    if language:
        label = paragraph.add_run(f"[{language}]\n")
        label.italic = True
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def _add_math_block(document: Document, lines: Iterable[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("\n".join(lines).strip())
    run.font.name = "Cambria Math"


def _add_list_paragraph(document: Document, text: str, *, ordered: bool, indent_chars: int) -> None:
    style = "List Number" if ordered else "List Bullet"
    paragraph = document.add_paragraph(style=style if _style_exists(document, style) else None)
    depth = min(6, max(0, indent_chars // 2))
    if depth:
        paragraph.paragraph_format.left_indent = Inches(0.25 * depth)

    if text.startswith("[ ] "):
        text = "☐ " + text[4:]
    elif text.lower().startswith("[x] "):
        text = "☒ " + text[4:]
    _add_inline(paragraph, text)


def markdown_to_docx(
    markdown_path: str | Path,
    output_path: str | Path | None = None,
    *,
    title: str | None = None,
    page_breaks: bool = True,
) -> Path:
    """Convert sanitizer-style Markdown into a standalone DOCX file.

    The converter intentionally preserves unsupported semantics rather than discarding them:
    LaTeX display math remains readable LaTeX in Cambria Math and Mermaid remains a code block.
    """

    source = Path(markdown_path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if source.suffix.lower() not in {".md", ".markdown"}:
        raise ValueError(f"Expected a Markdown file: {source}")
    output = (
        Path(output_path).expanduser().resolve()
        if output_path is not None
        else source.with_suffix(".docx")
    )
    output.parent.mkdir(parents=True, exist_ok=True)

    text = source.read_text(encoding="utf-8")
    # Sanitizer output should already be clean, but tolerate older generated files.
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    if title:
        document.add_heading(title, level=0)

    index = 0
    seen_page_marker = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        page_match = _PAGE_MARKER_RE.match(stripped)
        if page_match:
            if page_breaks and seen_page_marker:
                document.add_page_break()
            seen_page_marker = True
            index += 1
            continue

        fence_match = _FENCE_RE.match(line)
        if fence_match:
            marker = fence_match.group(1)
            language = fence_match.group(2)
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].lstrip().startswith(marker[:3]):
                block.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            _add_code_block(document, block, language)
            continue

        if stripped == "$$":
            math_lines: list[str] = []
            index += 1
            while index < len(lines) and lines[index].strip() != "$$":
                math_lines.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            _add_math_block(document, math_lines)
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            level = min(6, len(heading.group(1)))
            paragraph = document.add_heading(level=level)
            _add_inline(paragraph, heading.group(2))
            index += 1
            continue

        if _looks_like_table(lines, index):
            header = _split_table_row(lines[index])
            rows: list[list[str]] = []
            index += 2
            while index < len(lines):
                candidate = lines[index]
                if not candidate.strip() or "|" not in candidate:
                    break
                rows.append(_split_table_row(candidate))
                index += 1
            width = max([len(header), *(len(row) for row in rows)] or [1])
            table = document.add_table(rows=1, cols=width)
            if _style_exists(document, "Table Grid"):
                table.style = "Table Grid"
            for column, value in enumerate(header):
                cell = table.rows[0].cells[column]
                paragraph = cell.paragraphs[0]
                _add_inline(paragraph, value)
                for run in paragraph.runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for column in range(width):
                    value = row[column] if column < len(row) else ""
                    _add_inline(cells[column].paragraphs[0], value)
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            _add_list_paragraph(
                document,
                bullet.group(2).strip(),
                ordered=False,
                indent_chars=len(bullet.group(1).expandtabs(2)),
            )
            index += 1
            continue

        ordered = _ORDERED_RE.match(line)
        if ordered:
            _add_list_paragraph(
                document,
                ordered.group(2).strip(),
                ordered=True,
                indent_chars=len(ordered.group(1).expandtabs(2)),
            )
            index += 1
            continue

        quote = _BLOCKQUOTE_RE.match(line)
        if quote:
            paragraph = document.add_paragraph()
            if _style_exists(document, "Intense Quote"):
                paragraph.style = "Intense Quote"
            _add_inline(paragraph, quote.group(1))
            index += 1
            continue

        if re.fullmatch(r"\s*(?:-{3,}|\*{3,}|_{3,})\s*", line):
            paragraph = document.add_paragraph("────────────────────────")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue

        paragraph_text, next_index = _consume_paragraph(lines, index)
        paragraph = document.add_paragraph()
        _add_inline(paragraph, paragraph_text)
        index = max(index + 1, next_index)

    document.save(output)
    return output
