from pathlib import Path
from zipfile import ZipFile

import pymupdf
from docx import Document
from lxml import etree

from pdf_sanitizer.docx_export import markdown_to_docx


_W_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
_M_NS = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}


def _document_xml(path: Path):
    with ZipFile(path) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def _page_break_count(path: Path) -> int:
    root = _document_xml(path)
    return len(root.xpath('.//w:br[@w:type="page"]', namespaces=_W_NS))


def _math_count(path: Path) -> int:
    root = _document_xml(path)
    return len(root.xpath(".//m:oMath", namespaces=_M_NS))


def test_markdown_to_docx_preserves_core_structure(tmp_path: Path):
    markdown = tmp_path / "sample.md"
    markdown.write_text(
        """<!-- page: 1 -->

# Title

A **bold** paragraph with *italics* and `code`.

- [x] Finished
- [ ] Pending

| Name | Value |
| --- | --- |
| Alpha | 1 |

$$
x = y + 2
$$

```mermaid
flowchart LR
A --- B
```

<!-- page: 2 -->

## Second page

More text.
""",
        encoding="utf-8",
    )

    output = markdown_to_docx(markdown)
    assert output.suffix == ".docx"
    assert output.is_file()

    document = Document(output)
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "Title" in texts
    assert any("A bold paragraph" in text for text in texts)
    assert any("☒ Finished" in text for text in texts)
    assert any("☐ Pending" in text for text in texts)
    assert any("flowchart LR" in text for text in texts)
    assert "Second page" in texts
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "Name"
    assert document.tables[0].cell(1, 1).text == "1"
    assert _math_count(output) == 1
    # Word is reflowing by default; source PDF page comments are metadata.
    assert _page_break_count(output) == 0


def test_latex_math_becomes_native_word_omml_structures(tmp_path: Path):
    markdown = tmp_path / "math.md"
    markdown.write_text(
        r"""$$
\frac{1}{2}x^2 + \sqrt{y} + \alpha_1
$$
""",
        encoding="utf-8",
    )
    output = markdown_to_docx(markdown)
    root = _document_xml(output)
    assert len(root.xpath(".//m:oMath", namespaces=_M_NS)) == 1
    assert len(root.xpath(".//m:f", namespaces=_M_NS)) == 1
    assert len(root.xpath(".//m:rad", namespaces=_M_NS)) == 1
    assert len(root.xpath(".//m:sSup", namespaces=_M_NS)) >= 1
    assert len(root.xpath(".//m:sSub", namespaces=_M_NS)) >= 1


def test_source_page_breaks_are_explicit_opt_in(tmp_path: Path):
    markdown = tmp_path / "pages.md"
    markdown.write_text(
        "<!-- page: 1 -->\n\nFirst.\n\n<!-- page: 2 -->\n\nSecond.",
        encoding="utf-8",
    )
    output = markdown_to_docx(markdown, page_breaks=True)
    assert _page_break_count(output) == 1


def test_old_br_tags_are_flattened_during_docx_export(tmp_path: Path):
    markdown = tmp_path / "legacy.md"
    markdown.write_text("REVENUE<br>MANAGEMENT", encoding="utf-8")
    output = markdown_to_docx(markdown)
    document = Document(output)
    assert document.paragraphs[0].text == "REVENUE MANAGEMENT"


def test_visual_placeholder_becomes_readable_caption_without_bbox(tmp_path: Path):
    markdown = tmp_path / "visual.md"
    markdown.write_text(
        '[IMAGE_PLACEHOLDER page=62 bbox="59,347,396,370"]',
        encoding="utf-8",
    )
    output = markdown_to_docx(markdown)
    paragraph = Document(output).paragraphs[0]
    assert paragraph.style.name == "Caption"
    assert paragraph.text == "Image omitted from semantic extraction (source page 62)."
    assert "bbox" not in paragraph.text


def test_source_pdf_can_embed_visual_placeholder_crop(tmp_path: Path):
    source_pdf = tmp_path / "source.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page(width=300, height=200)
    page.draw_rect((50, 50, 250, 120), color=(0, 0, 0), width=2)
    page.insert_text((70, 90), "visual formula")
    pdf.save(source_pdf)
    pdf.close()

    markdown = tmp_path / "visual.md"
    markdown.write_text(
        '<!-- page: 1 -->\n\n[IMAGE_PLACEHOLDER page=1 bbox="50,50,250,120"]',
        encoding="utf-8",
    )
    output = markdown_to_docx(markdown, source_pdf=source_pdf)
    root = _document_xml(output)
    with ZipFile(output) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert media
    assert len(root.xpath(".//w:drawing", namespaces=_W_NS)) == 1
    assert "omitted from semantic extraction" not in " ".join(
        paragraph.text for paragraph in Document(output).paragraphs
    )


def test_html_inline_semantics_become_word_run_formatting(tmp_path: Path):
    markdown = tmp_path / "inline.md"
    # The unmatched '*' before <sup> is deliberate. Older italic matching could swallow
    # the tag and leave literal HTML in the generated Word document.
    markdown.write_text(
        'J(p*) = value <sup>2</sup>, <sub>t</sub>, and <u>important</u>.',
        encoding="utf-8",
    )
    output = markdown_to_docx(markdown)
    paragraph = Document(output).paragraphs[0]
    assert "<sup>" not in paragraph.text
    assert "<sub>" not in paragraph.text
    assert "<u>" not in paragraph.text
    assert any(run.text == "2" and run.font.superscript for run in paragraph.runs)
    assert any(run.text == "t" and run.font.subscript for run in paragraph.runs)
    assert any(run.text == "important" and run.underline for run in paragraph.runs)


def test_inline_source_page_marker_is_invisible_in_word_text(tmp_path: Path):
    markdown = tmp_path / "joined.md"
    markdown.write_text(
        "This is signifi<!-- page: 2 -->cant evidence.",
        encoding="utf-8",
    )
    output = markdown_to_docx(markdown)
    assert Document(output).paragraphs[0].text == "This is significant evidence."
