from pathlib import Path

from docx import Document

from pdf_sanitizer.docx_export import markdown_to_docx


def test_markdown_to_docx_preserves_core_structure(tmp_path: Path):
    markdown = tmp_path / "sample.md"
    markdown.write_text(
        """<!-- page: 1 -->

# Title

A **bold** paragraph with *italics* and `code`.

- [x] Finished
- [ ] Pending

| Name | Value |
| --- | --- |
| Alpha | 1 |

$$
x = y + 2
$$

```mermaid
flowchart LR
A --- B
```

<!-- page: 2 -->

## Second page

More text.
""",
        encoding="utf-8",
    )

    output = markdown_to_docx(markdown)
    assert output.suffix == ".docx"
    assert output.is_file()

    document = Document(output)
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "Title" in texts
    assert any("A bold paragraph" in text for text in texts)
    assert any("☒ Finished" in text for text in texts)
    assert any("☐ Pending" in text for text in texts)
    assert any("x = y + 2" in text for text in texts)
    assert any("flowchart LR" in text for text in texts)
    assert "Second page" in texts
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "Name"
    assert document.tables[0].cell(1, 1).text == "1"


def test_old_br_tags_are_flattened_during_docx_export(tmp_path: Path):
    markdown = tmp_path / "legacy.md"
    markdown.write_text("REVENUE<br>MANAGEMENT", encoding="utf-8")
    output = markdown_to_docx(markdown)
    document = Document(output)
    assert document.paragraphs[0].text == "REVENUE MANAGEMENT"
